import marimo

__generated_with = "0.23.9"
app = marimo.App(width="medium")


@app.cell
def _():
    import marimo as mo

    from datetime import datetime
    import re

    import docx
    import duckdb
    import numpy
    import polars as pl

    return datetime, docx, duckdb, mo, pl, re


@app.cell
def _(mo):
    root_dir = mo.notebook_dir() / "S13"
    if not root_dir.exists():
        raise Exception(f"Le répertoire S13 ({root_dir}) n'existe pas.")

    villes = [
      "Andilly",
      "Montmagny",
      "Eaubonne",
      "Montmorency",
      "Margency",
      "Groslay",
      "Deuil",
      "Soisy sous Montmorency"
    ]
    return root_dir, villes


@app.function
def document_to_rows(document):
    rows = []
    for table in document.tables:
        rows.extend([
            [cell.text.strip() for cell in row.cells]
            for row in  table.rows
        ])
    return rows


@app.cell
def _(pl):
    def rows_to_df(rows):
        schema = ["a","b","c","d","e","f","g","h","i","j"]
        return pl.DataFrame(rows, schema=schema, orient="row")

    return (rows_to_df,)


@app.cell
def _(datetime, docx, duckdb, mo, re, root_dir, rows_to_df, villes):
    def parse_date(x: str) -> duckdb.sqltypes.DATE:
        if not x:
            return
        cleaned = re.sub(r"[^0-9]+", "/", x)
        try:
            return datetime.strptime(cleaned, "%d/%m/%Y").date()
        except:
            try:
                return datetime.strptime(cleaned, "%d/%m/%y").date()
            except:
                raise Exception(f"⚠ Format de date incorrect : {ville.name} / {fichier.name} / {x}")

    try:
        duckdb.remove_function("parse_date")
    except:
        pass
    _ = duckdb.create_function("parse_date", parse_date, null_handling="special")

    for _ville in mo.status.progress_bar(villes):
        ville = root_dir / _ville
        for fichier in ville.iterdir():
            if mo.app_meta().mode == "script":
                print(f"Chargement du fichier {fichier}")
            
            document = docx.Document(fichier)

            rows = document_to_rows(document)

            raw_df = rows_to_df(rows)

            odd_df = mo.sql(
                f"""
                select *
                  from raw_df
                qualify row_number() over () % 2 = 1
                offset 1
                """, output=False
            )

            even_df = mo.sql(
                f"""
                select *
                  from raw_df
                qualify row_number() over () % 2 = 0
                offset 1
                """, output=False
            )

            parsed_df = mo.sql(
                f"""
                select '{ville.name}' ville, '{fichier.name}' fichier,
                       odd_df."a" territoire,
                       parse_date(coalesce(odd_df."b", even_df."b")) derniere_fois_le,
                       coalesce(odd_df."c", odd_df."d") attribue_a1,
                       coalesce(odd_df."e", odd_df."f") attribue_a2,
                       coalesce(odd_df."g", odd_df."h") attribue_a3,
                       coalesce(odd_df."i", odd_df."j") attribue_a4,
                       parse_date(even_df."c") attribue_le1,
                       parse_date(even_df."d") entierement_parcouru_le1,
                       parse_date(even_df."e") attribue_le2,
                       parse_date(even_df."f") entierement_parcouru_le2,
                       parse_date(even_df."g") attribue_le3,
                       parse_date(even_df."h") entierement_parcouru_le3,
                       parse_date(even_df."i") attribue_le4,
                       parse_date(even_df."j") entierement_parcouru_le4
                  from odd_df
                  left join even_df
                    on odd_df."a" = even_df."a"
                """, output=False
            )
            _query = "select distinct * from parsed_df where coalesce(territoire, '') <> ''"
            try:
                mo.sql(f"create table territoires_complet as {_query}", output=False)
            except:
                mo.sql(f"insert into territoires_complet {_query}", output=False)
    mo.sql("select * from territoires_complet")
    return even_df, odd_df, raw_df


@app.cell(hide_code=True)
def _(mo, territoires_complet):
    _df = mo.sql(
        f"""
        drop table if exists territoire
        ;
        drop table if exists attribution
        ;
        create table territoire as
        select distinct ville, fichier, territoire, derniere_fois_le
          from territoires_complet
        ;
        create table attribution as
        select distinct ville, fichier, territoire,
                        attribue_a1 attribue_a, attribue_le1 attribue_le, entierement_parcouru_le1 entierement_parcouru_le
          from territoires_complet
        ;
        insert into attribution
        select distinct ville, fichier, territoire,
                        attribue_a2 attribue_a, attribue_le2 attribue_le, entierement_parcouru_le2 entierement_parcouru_le
          from territoires_complet
        ;
        insert into attribution
        select distinct ville, fichier, territoire,
                        attribue_a3 attribue_a, attribue_le3 attribue_le, entierement_parcouru_le3 entierement_parcouru_le
          from territoires_complet
        ;
        insert into attribution
        select distinct ville, fichier, territoire,
                        attribue_a4 attribue_a, attribue_le4 attribue_le, entierement_parcouru_le4 entierement_parcouru_le
          from territoires_complet
        ;
        alter table territoire add column sorti bool
        ;
        update territoire
           set sorti = exists (select 1
                                 from attribution
                                where attribution.ville = territoire.ville
                                  and attribution.fichier = territoire.fichier
                                  and attribution.territoire = territoire.territoire
                                  and attribution.attribue_le is not null
                                  and attribution.entierement_parcouru_le is null)
        ;
        table territoire
        ;
        -- table attribution
        -- ;
        """
    )
    return


@app.cell(hide_code=True)
def _(attribution, mo, territoire):
    analyse_s13_df = mo.sql(
        f"""
        with t as (
        select distinct attribution.ville,
                        attribution.fichier,
                        attribution.territoire,
                        attribution.attribue_a,
                        attribution.entierement_parcouru_le,
                        if(territoire.sorti, 'X', '') sorti--.strftime('%d/%m/%y') entierement_parcouru_le
             --, row_number() over (partition by ville, fichier, territoire order by entierement_parcouru_le desc nulls last)
          from attribution
          join territoire
            on territoire.ville = attribution.ville
           and territoire.fichier = attribution.fichier
           and territoire.territoire = attribution.territoire
           --and territoire.sorti = False
        qualify row_number() over (partition by attribution.ville, attribution.fichier, attribution.territoire order by attribution.entierement_parcouru_le desc nulls last) = 1
         order by territoire.sorti, attribution.ville, attribution.entierement_parcouru_le nulls first, attribution.fichier, attribution.territoire -- on fait le order ici, parce qu'à l'étape suivante entierement_parcouru_le devient un str
        )
        select ville, fichier, territoire, attribue_a, entierement_parcouru_le.strftime('%d/%m/%y') entierement_parcouru_le, sorti
          from t
        """
    )
    return (analyse_s13_df,)


@app.cell
def _(analyse_s13_df, mo):
    _xlsx_output_path = mo.notebook_dir() / "analyse_s13.xlsx"
    analyse_s13_df.write_excel(_xlsx_output_path)

    print(f"Fichier {_xlsx_output_path} créé.")
    return


if __name__ == "__main__":
    app.run()
